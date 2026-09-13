"""Step 0 item 7 (read-only): flag / opt-out / LID values on the in-scope contacts,
plus HR-value length checks against the describe lengths."""
import os, json, subprocess, sys, collections
import openpyxl
from openpyxl.utils import column_index_from_string as ci

for v in ("HTTP_PROXY", "HTTPS_PROXY", "http_proxy", "https_proxy"):
    os.environ.pop(v, None)
ORG = sys.argv[1] if len(sys.argv) > 1 else "FULLUAT"
MAP = r"C:\Users\Kamyar.Jannati\Downloads\Centellic_Employee_to_Salesforce_Mapping_2026-09-07.xlsx"
HR = r"C:\Users\Kamyar.Jannati\Downloads\Current_Employee_List_non_confidential_7.9.26.xlsx"  # copy of the OneDrive original; the OneDrive path is not readable from the sandboxed interpreter

wb = openpyxl.load_workbook(MAP, read_only=True, data_only=True)
emp = list(wb["Employees"].iter_rows(values_only=True))
c = lambda L: ci(L) - 1
contacts = {str(r[c('AF')]).strip()[:18]: r[c('C')] for r in emp[1:] if r[c('AF')]}
users = {str(r[c('AD')]).strip()[:18]: r[c('C')] for r in emp[1:] if r[c('AD')] and str(r[c('AE')]).strip().lower() == 'yes'}

# HR lengths
hw = openpyxl.load_workbook(HR, read_only=True, data_only=True)["Sheet1"]
hrows = list(hw.iter_rows(min_row=2, values_only=True))
hh = [str(x).strip() if x else "" for x in hrows[0]]
iEmp, iTitle, iName = hh.index("Employee ID"), hh.index("Business Title"), hh.index("Preferred Name")
ids = [str(r[iEmp]).strip() for r in hrows[1:] if r[iEmp] is not None]
titles = [str(r[iTitle]).strip() for r in hrows[1:] if r[iTitle]]
print("HR rows:", len(hrows) - 1)
print("Employee ID: max len %d, min len %d, with leading zero %d, non-numeric %d, sample %s"
      % (max(map(len, ids)), min(map(len, ids)), sum(1 for x in ids if x.startswith('0')), sum(1 for x in ids if not x.isdigit()), ids[:3]))
print("Employee ID cell types in workbook:", collections.Counter(type(r[iEmp]).__name__ for r in hrows[1:] if r[iEmp] is not None))
print("Business Title: max len %d, >80 chars %d (User.Title), >128 chars %d (Contact.Title)" % (max(map(len, titles)), sum(1 for t in titles if len(t) > 80), sum(1 for t in titles if len(t) > 128)))
for t in sorted({t for t in titles if len(t) > 80}, key=len, reverse=True)[:10]: print("   >80:", len(t), t)

def query(soql):
    o = subprocess.run(["sf", "data", "query", "--target-org", ORG, "--query", soql, "--json"], capture_output=True, text=True, shell=True, encoding="utf-8")
    j = json.loads(o.stdout)
    if j.get("status") != 0: print(o.stdout[:1500]); sys.exit(1)
    return j["result"]["records"]

recs = {}
idl = sorted(contacts)
for i in range(0, len(idl), 200):
    soql = ("SELECT Id, Name, Title, Employee_Number__c, Not_at_Company_Checkbox__c, No_Longer_w_Company__c, ADvendio__LeftCompany__c, "
            "LID__No_longer_at_Company__c, Not_at_Company_Flag_Status__c, DoNotCall, HasOptedOutOfEmail, No_Postal__c, ReportsToId, Account.Name "
            "FROM Contact WHERE Id IN (%s)" % ",".join("'%s'" % x for x in idl[i:i+200]))
    for r in query(soql): recs[r["Id"][:18]] = r
print("\n=== %s: %d of %d in-scope contacts found" % (ORG, len(recs), len(contacts)))
flagged = [r for r in recs.values() if r["Not_at_Company_Checkbox__c"] or r["No_Longer_w_Company__c"] or r["ADvendio__LeftCompany__c"]]
print("flag counts: Not_at_Company=%d  No_Longer_w_Company=%d  ADvendio_LeftCompany=%d  any=%d" % (
    sum(1 for r in recs.values() if r["Not_at_Company_Checkbox__c"]), sum(1 for r in recs.values() if r["No_Longer_w_Company__c"]),
    sum(1 for r in recs.values() if r["ADvendio__LeftCompany__c"]), len(flagged)))
print("LID__No_longer_at_Company__c values across all found:", collections.Counter(r["LID__No_longer_at_Company__c"] for r in recs.values()))
print("opt-outs across all found: DoNotCall=%d HasOptedOutOfEmail=%d No_Postal=%d" % (
    sum(1 for r in recs.values() if r["DoNotCall"]), sum(1 for r in recs.values() if r["HasOptedOutOfEmail"]), sum(1 for r in recs.values() if r["No_Postal__c"])))
print("\nflagged contacts (D4/D5):")
print("%-18s %-28s NAC NLW ADV  DNC EOO NoPost  LID / status                 account" % ("Id", "Name"))
for r in flagged:
    print("%-18s %-28s %-3s %-3s %-3s  %-3s %-3s %-6s  %s / %s   %s" % (r["Id"][:18], r["Name"][:28],
          int(r["Not_at_Company_Checkbox__c"]), int(r["No_Longer_w_Company__c"]), int(r["ADvendio__LeftCompany__c"]),
          int(r["DoNotCall"]), int(r["HasOptedOutOfEmail"]), int(r["No_Postal__c"]), r["LID__No_longer_at_Company__c"], r["Not_at_Company_Flag_Status__c"],
          (r["Account"] or {}).get("Name")))
print("\nEmployee_Number__c current values: blank %d, filled %d, sample %s" % (
    sum(1 for r in recs.values() if not r["Employee_Number__c"]), sum(1 for r in recs.values() if r["Employee_Number__c"]),
    [r["Employee_Number__c"] for r in recs.values() if r["Employee_Number__c"]][:5]))
json.dump({"org": ORG, "flagged": flagged}, open(os.path.join(os.path.dirname(__file__), "item7_%s.json" % ORG), "w"), indent=1)

# User.EmployeeNumber legacy values, for the tech-debt register and the report grep
urecs = {}
ul = sorted(users)
for i in range(0, len(ul), 200):
    for r in query("SELECT Id, Name, EmployeeNumber, Employee_Code__c, Title FROM User WHERE Id IN (%s)" % ",".join("'%s'" % x for x in ul[i:i+200])):
        urecs[r["Id"][:18]] = r
print("\n=== %s: %d of %d in-scope users found" % (ORG, len(urecs), len(users)))
en = [r["EmployeeNumber"] for r in urecs.values() if r["EmployeeNumber"]]
print("User.EmployeeNumber filled %d, blank %d; non-numeric values (legacy codes): %s" % (len(en), len(urecs) - len(en), sorted({x for x in en if not x.isdigit()})[:40]))
ec = [r["Employee_Code__c"] for r in urecs.values() if r["Employee_Code__c"]]
print("User.Employee_Code__c filled %d, blank %d; non-numeric: %s" % (len(ec), len(urecs) - len(ec), sorted({x for x in ec if not x.isdigit()})[:40]))
