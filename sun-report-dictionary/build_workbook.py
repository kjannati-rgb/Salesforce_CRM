# Phase 4 - build out/SUN_Report_Data_Dictionary.xlsx from out/dictionary_data.json
import json, os, datetime, subprocess
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

BASE = os.path.dirname(os.path.abspath(__file__))
data = json.load(open(os.path.join(BASE, 'out/dictionary_data.json'), encoding='utf-8'))
meta = data['meta']

HDR_FILL = PatternFill('solid', fgColor='1F3864')
HDR_FONT = Font(name='Arial', bold=True, color='FFFFFF')
BODY_FONT = Font(name='Arial')
WRAP = Alignment(wrap_text=True, vertical='top')
TOP = Alignment(vertical='top')

def style_table(ws, headers, widths, wrap_cols=()):
    # headers must already be the sheet's first appended row; this styles them
    for ci in range(1, len(headers) + 1):
        c = ws.cell(row=1, column=ci)
        c.fill = HDR_FILL; c.font = HDR_FONT; c.alignment = Alignment(vertical='center')
    for ci, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(ci)].width = w
    ws.freeze_panes = 'A2'
    for row in ws.iter_rows(min_row=2):
        for c in row:
            if c.value is not None:
                c.font = BODY_FONT
                c.alignment = WRAP if c.column in wrap_cols else TOP

wb = Workbook()

# ---------- Sheet 1: Overview ----------
ws = wb.active
ws.title = 'Overview'
try:
    cli_ver = subprocess.run(['sf', '--version'], capture_output=True, text=True, shell=True).stdout.strip().split('\n')[0]
except Exception:
    cli_ver = 'unknown'
sdf = meta['standardDateFilter'] or {}
rows = [
    ('Report Name', 'SUN Report'),
    ('Report ID', '00O6g000005mExIEAU'),
    ('Report DeveloperName', 'SUN_Report_ZM_HoN'),
    ('Folder DeveloperName', 'FinanceReports'),
    ('Org', 'Production, Org ID 00D6g0000081IOg (lawbusinessresearch.my.salesforce.com)'),
    ('Authenticated user', 'kamyar.jannati@lbresearch.com'),
    ('Format', 'Tabular'),
    ('Report Type', '%s (%s)' % (meta['reportType']['label'], meta['reportType']['type'])),
    ('Report Type kind', 'Custom Report Type - base Opportunity, joins OpportunityLineItems then OpportunityLineItemSchedules (see Report Type Map sheet)'),
    ('Known purpose', 'Finance extract to SunSystems. Closed Won opportunities only, excludes 0-value lines'),
    ('Scope', meta['scope'] or ''),
    ('Standard date filter', '%s %s (current window %s to %s)' % (sdf.get('column',''), sdf.get('durationValue',''), sdf.get('startDate',''), sdf.get('endDate',''))),
    ('Report currency', meta['currency'] or ''),
    ('Boolean filter logic', meta['booleanFilter'] or 'None - all filters AND'),
    ('Column count', len(data['columns'])),
    ('Filter count', len(data['filters'])),
    ('Row-level formulas', len(meta['cdf'])),
    ('API version used', 'v64.0'),
    ('Generated', datetime.datetime.now().strftime('%Y-%m-%d %H:%M')),
    ('sf CLI version', cli_ver),
    ('', ''),
    ('Audit trail', ''),
    ('raw/analytics_describe.json', 'Analytics REST describe - column labels, data types, filters, row-level formula. Feeds Columns + Filters sheets'),
    ('raw/SUN_Report_ZM_HoN.report-meta.xml', 'Metadata API report XML - canonical stored definition. Cross-checked against Analytics describe (80 columns, 6 filters, both agree)'),
    ('raw/Opportunities_With_or_Without_Products_Schedules.reportType-meta.xml', 'Custom Report Type XML - join structure. Feeds Report Type Map sheet'),
    ('describes/*.json', '11 object describes - field labels, types, lengths, formulas, picklists. Feeds Columns, Filters, Picklists sheets'),
]
for ri, (k, v) in enumerate(rows, 1):
    a = ws.cell(row=ri, column=1, value=k); a.font = Font(name='Arial', bold=True)
    b = ws.cell(row=ri, column=2, value=v); b.font = BODY_FONT; b.alignment = WRAP
ws.column_dimensions['A'].width = 45
ws.column_dimensions['B'].width = 110

# ---------- Sheet 2: Columns ----------
ws = wb.create_sheet('Columns')
headers = ['#', 'Report Column Label', 'Raw Token', 'Source Object', 'Field API Name', 'Field Label',
           'Data Type', 'Length / Precision.Scale', 'Custom', 'Formula?', 'Formula Text',
           'Picklist?', 'Picklist Values', 'Lookup To', 'Help Text', 'Notes']
ws.append(headers)
for r in data['columns']:
    ws.append([r['n'], r['label'], r['token'], r['source_object'], r['field_api'], r['field_label'],
               r['data_type'], r['len_prec'], r['custom'], r['formula'], r['formula_text'],
               r['picklist'], r['picklist_values'], r['lookup_to'], r['help_text'], r['notes']])
style_table(ws, headers,
            [5, 30, 42, 22, 30, 28, 12, 10, 8, 9, 60, 9, 40, 14, 45, 45],
            wrap_cols=(11, 13, 15, 16))
ws.auto_filter.ref = 'A1:P%d' % ws.max_row

# ---------- Sheet 3: Filters ----------
ws = wb.create_sheet('Filters')
headers = ['#', 'Raw Token', 'Source Object', 'Field API Name', 'Operator', 'Value(s)', 'Notes']
ws.append(headers)
for r in data['filters']:
    ws.append([r['n'], r['token'], r['source_object'], r['field_api'], r['operator'], r['value'], r['notes']])
style_table(ws, headers, [5, 55, 24, 32, 12, 45, 45], wrap_cols=(6, 7))
ws.auto_filter.ref = 'A1:G%d' % ws.max_row
foot = ws.max_row + 2
sdf = meta['standardDateFilter'] or {}
footer = [
    ('Boolean filter logic', meta['booleanFilter'] or 'None - all 6 filters combined with AND'),
    ('Scope', meta['scope'] or ''),
    ('Standard date filter', '%s = %s (window at generation time: %s to %s)' % (sdf.get('column',''), sdf.get('durationValue',''), sdf.get('startDate',''), sdf.get('endDate',''))),
    ('Report currency', meta['currency'] or ''),
    ('Cross-check', 'XML criteriaItems and Analytics reportFilters agree 1:1 on column, operator and values. One notation difference only: XML stores TOTAL_GROSS__c filter value as "0", Analytics renders it "0.00". Same filter'),
]
for k, v in footer:
    a = ws.cell(row=foot, column=1, value=k); a.font = Font(name='Arial', bold=True)
    b = ws.cell(row=foot, column=2, value=v); b.font = BODY_FONT
    foot += 1

# ---------- Sheet 4: Report Type Map ----------
ws = wb.create_sheet('Report Type Map')
headers = ['Level', 'Object / Join', 'Join semantics (XML)', 'Columns used from this join']
used = {}
for r in data['columns']:
    tok = r['token']
    if tok.startswith('OpportunityLineItemSchedule'):
        j = 'Opportunity.OpportunityLineItems.OpportunityLineItemSchedules'
    elif tok.startswith('OpportunityLineItem'):
        j = 'Opportunity.OpportunityLineItems'
    elif tok in data['meta']['cdf']:
        j = 'Report-defined'
    else:
        j = 'Opportunity'
    used[j] = used.get(j, 0) + 1
rows = [
    ('Base', 'Opportunity', 'n/a (base object)', used.get('Opportunity', 0)),
    ('Join 1', 'OpportunityLineItems (OpportunityLineItem)', 'outerJoin=false in CRT XML = "with" related records. NOTE: CRT label says "With or Without" but stored XML is an inner join. Discrepancy flagged', used.get('Opportunity.OpportunityLineItems', 0)),
    ('Join 2', 'OpportunityLineItemSchedules (OpportunityLineItemSchedule)', 'outerJoin=false in CRT XML = "with" related records. Same label/XML discrepancy as Join 1', used.get('Opportunity.OpportunityLineItems.OpportunityLineItemSchedules', 0)),
    ('n/a', 'Report-defined (row-level formula CDF1)', 'n/a', used.get('Report-defined', 0)),
]
note = ('Columns counted by token prefix. Lookup traversals (Invoice_Contact__c, Account, Owner, Order__c, '
        'PricebookEntry.Product2, SBQQ__QuoteLine__c, SBQQ__Quote__c, Bio_Taker__c) are reached through these '
        'three tables via lookup, not separate CRT joins.')
ws.append(headers)
for r in rows: ws.append(list(r))
style_table(ws, headers, [8, 52, 70, 28], wrap_cols=(3,))
ws.cell(row=ws.max_row + 2, column=1, value=note).font = BODY_FONT
ws.cell(row=ws.max_row, column=1).alignment = WRAP
ws.merge_cells(start_row=ws.max_row, start_column=1, end_row=ws.max_row, end_column=4)

# ---------- Sheet 5: Picklists ----------
if data['picklists']:
    ws = wb.create_sheet('Picklists')
    headers = ['Object', 'Field API Name', 'Value', 'Active', 'Default']
    ws.append(headers)
    for r in data['picklists']:
        ws.append(list(r))
    style_table(ws, headers, [24, 32, 55, 9, 9])
    ws.auto_filter.ref = 'A1:E%d' % ws.max_row

out_path = os.path.join(BASE, 'out/SUN_Report_Data_Dictionary.xlsx')
wb.save(out_path)
print('saved', out_path)
print('picklist sheet rows:', len(data['picklists']))
